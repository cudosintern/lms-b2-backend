"""LMS consolidated marks, using the migrated curriculum/course/section tables."""
from collections import defaultdict
from io import BytesIO, StringIO
import csv

from fastapi import HTTPException
from sqlalchemy import bindparam, text


def _rows(db, sql, params=None, expanding=()):
    statement = text(sql)
    for key in expanding:
        statement = statement.bindparams(bindparam(key, expanding=True))
    return [dict(row) for row in db.execute(statement, params or {}).mappings()]


def get_department_options(db, org_id):
    return _rows(db, """SELECT dept_id AS id, dept_name AS name FROM iems_department
        WHERE (:org_id IS NULL OR org_id = :org_id OR org_id IS NULL)
        ORDER BY dept_name""", {"org_id": org_id})


def get_curriculum_options(db, org_id, department_id=None):
    return _rows(db, """SELECT academic_batch_id, academic_batch_id AS crclm_id,
        COALESCE(academic_batch_desc, academic_batch_code) AS name, dept_id, pgm_id
        FROM iems_academic_batch
        WHERE (:org_id IS NULL OR org_id = :org_id OR org_id IS NULL)
          AND (:department_id IS NULL OR dept_id = :department_id)
        ORDER BY name""", {"org_id": org_id, "department_id": department_id})


def get_term_options(db, academic_batch_id):
    rows = _rows(db, """SELECT semester_id, semester_id AS crclm_term_id,
        semester AS semester_number, semester_desc AS name FROM iems_semester
        WHERE academic_batch_id = :batch AND semester IS NOT NULL
        ORDER BY semester, semester_id""", {"batch": academic_batch_id})
    for row in rows:
        row["name"] = row["name"] or f"Semester {row['semester_number']}"
    return rows


def _term(db, academic_batch_id, semester_id=None, crclm_term_id=None):
    # crclm_term_id is retained only as an API alias for the new semester PK.
    if semester_id is not None and crclm_term_id is not None and semester_id != crclm_term_id:
        raise HTTPException(422, "semester_id and crclm_term_id must identify the same term")
    term_id = semester_id if semester_id is not None else crclm_term_id
    if term_id is None:
        raise HTTPException(422, "Select a term")
    rows = _rows(db, """SELECT semester_id, semester, semester_desc FROM iems_semester
        WHERE academic_batch_id = :batch AND semester_id = :term""",
        {"batch": academic_batch_id, "term": term_id})
    if not rows or rows[0]["semester"] is None:
        raise HTTPException(404, "Term does not belong to the selected curriculum")
    return rows[0]


MAPPED_COURSES = """
    FROM cudos_map_courseto_course_instructor m
    JOIN iems_courses c ON c.crs_id = m.crs_id
    JOIN iems_semester sem ON sem.semester_id = m.semester_id
        AND sem.academic_batch_id = m.academic_batch_id
    JOIN cudos_master_type_details sec ON sec.mt_details_id = m.section_id
    WHERE m.academic_batch_id = :batch AND m.semester_id = :term
      AND c.academic_batch_id = m.academic_batch_id AND c.semester = sem.semester
"""


def get_section_options(db, academic_batch_id, semester_id=None, crclm_term_id=None):
    term = _term(db, academic_batch_id, semester_id, crclm_term_id)
    return _rows(db, "SELECT DISTINCT sec.mt_details_id AS section_id, "
        "sec.mt_details_name AS section_name " + MAPPED_COURSES + " ORDER BY section_name",
        {"batch": academic_batch_id, "term": term["semester_id"]})


def get_course_options(db, academic_batch_id, semester_id=None, crclm_term_id=None, section_id=None):
    term = _term(db, academic_batch_id, semester_id, crclm_term_id)
    return _rows(db, """SELECT DISTINCT c.crs_id AS course_id, c.crs_code AS course_code,
        COALESCE(c.crs_title, c.crs_code) AS course_title, c.semester """ + MAPPED_COURSES +
        " AND (:section IS NULL OR m.section_id = :section) ORDER BY course_code, course_id",
        {"batch": academic_batch_id, "term": term["semester_id"], "section": section_id})


def _number(value):
    return None if value is None else round(float(value), 2)


def _ems_data(db, params):
    """Normalize EMS assessment marks without adding CIA aggregates twice."""
    registrations = _rows(db, """SELECT sc.*, s.usno AS student_usn, c.crs_id,
        c.cia_max_marks, c.ise_max_marks, c.mse_max_marks, c.see_max_marks,
        c.viva_max_marks, c.tw_max_marks
        FROM iems_student_courses sc
        JOIN iems_courses c ON c.crs_code = sc.crs_code AND c.academic_batch_id = sc.batch_id
            AND c.semester = sc.semester
        JOIN iems_students s ON s.academic_batch_id = sc.batch_id
            AND (s.usno = NULLIF(sc.usno, '') OR (s.regno = NULLIF(sc.regno, '')
                AND NOT EXISTS (SELECT 1 FROM iems_students identified
                    WHERE identified.academic_batch_id = sc.batch_id AND identified.usno = NULLIF(sc.usno, ''))))
        JOIN iems_semester sem ON sem.semester_id = :term AND sem.academic_batch_id = :batch
        WHERE sc.batch_id = :batch AND sc.semester = sem.semester AND c.crs_id IN :course_ids
            AND COALESCE(sc.is_withdrawn, 0) = 0 AND COALESCE(sc.is_drop, 0) = 0
        ORDER BY sc.result_year, sc.std_crs_id""", params, ("course_ids",))
    latest = {(r["student_usn"], r["crs_id"]): r for r in registrations}
    if not latest:
        return {}, [], {}
    details = _rows(db, """SELECT marks.id, marks.std_crs_id, marks.occasion_id,
        marks.cia_master_id, marks.secured_marks, marks.is_absentee, marks.result_year,
        occasion.cia_occasion_type_desc, occasion.cia_occasion_type_code,
        master.cia_master_name, master.cia_max_marks
        FROM iems_cia_student_courses marks
        LEFT JOIN iems_cia_occasion_type occasion ON occasion.cia_occasion_type_id = marks.occasion_id
        LEFT JOIN iems_cia_exam_master master ON master.id = marks.cia_master_id
        WHERE marks.std_crs_id IN :ids ORDER BY marks.result_year, marks.id""",
        {"ids": [r["std_crs_id"] for r in latest.values()]}, ("ids",))
    by_registration = defaultdict(dict)
    for d in details:
        key = f"cia:{d['cia_master_id']}:{d['occasion_id']}"
        by_registration[d["std_crs_id"]][key] = d
    definitions, marks, totals = {}, [], {}
    for identity, r in latest.items():
        components = []
        detailed = by_registration[r["std_crs_id"]]
        for key, d in detailed.items():
            label = d["cia_occasion_type_desc"] or d["cia_occasion_type_code"] or "CIA"
            if d["cia_master_name"]:
                label = f"{d['cia_master_name']} - {label}"
            components.append((key, label, d["cia_max_marks"], d["secured_marks"], bool(d["is_absentee"])))
        if not detailed:
            if r["total_cia"] is not None:
                components.append(("cia", "CIA Total", r["cia_max_marks"], r["total_cia"], False))
            else:
                for key in ("ise", "mse"):
                    if r[f"total_{key}"] is not None:
                        components.append((key, f"{key.upper()} Total", r[f"{key}_max_marks"], r[f"total_{key}"], False))
        see = r["see_actual"] if r["see_actual"] is not None else r["see"]
        tw = r["tw_marks_actual"] if r["tw_marks_actual"] is not None else r["tw_marks"]
        for key, label, value, absent in (("see", "SEE", see, r["see_absentee"]),
                ("viva", "Viva", r["viva_marks"], r["viva_absentee"]), ("tw", "TW", tw, r["tw_absentee"])):
            if value is not None or absent:
                components.append((key, label, r[f"{key}_max_marks"], value, bool(absent)))
        if not components and r["cia_see"] is not None:
            components.append(("combined", "CIA + SEE", (r["cia_max_marks"] or 0) + (r["see_max_marks"] or 0), r["cia_see"], False))
        for key, label, maximum, value, absent in components:
            definitions[(r["crs_id"], key)] = {"component_id": key, "occasion_name": label,
                "max_marks": _number(maximum), "marks": None, "status": "missing", "source": "ems"}
            # NULL without an explicit absence flag remains missing.
            if value is not None or absent:
                marks.append({"student_usn": r["student_usn"], "crs_id": r["crs_id"],
                    "qpd_id": key, "total_marks": None if absent else value})
        cia = r["total_cia"]
        if cia is None:
            internal = [r[k] for k in ("total_ise", "total_mse") if r[k] is not None]
            if not internal:
                internal = [d["secured_marks"] for d in detailed.values() if d["secured_marks"] is not None and not d["is_absentee"]]
            cia = sum(internal) if internal else None
        values = [v for v in (cia, None if r["see_absentee"] else see,
            None if r["viva_absentee"] else r["viva_marks"], None if r["tw_absentee"] else tw) if v is not None]
        totals[identity] = round(sum(values), 2) if values else (
            _number(r["cia_see"]) if any(c[0] == "combined" for c in components) else None)
    return definitions, marks, totals


def get_organisation_marks_source(db, org_id):
    organisations = _rows(db, "SELECT marks_fetch_from FROM iems_organisation WHERE org_id = :org_id", {"org_id": org_id})
    if not organisations:
        raise HTTPException(404, "Organisation not found")
    setting = organisations[0]["marks_fetch_from"]
    if setting is None:
        raise HTTPException(422, "Marks source is not configured for this organisation")
    # CodeIgniter uses 0 for LMS; a nonzero setting selects EMS.
    return "lms" if int(setting) == 0 else "ems"


def build_consolidated_student_marks_report(db, request, org_id=1):
    marks_source = get_organisation_marks_source(db, org_id)
    term = _term(db, request.academic_batch_id, request.semester_id, request.crclm_term_id)
    batches = _rows(db, """SELECT academic_batch_desc, academic_batch_code, dept_id
        FROM iems_academic_batch WHERE academic_batch_id = :batch""",
        {"batch": request.academic_batch_id})
    if not batches or (request.department_id is not None and batches[0]["dept_id"] != request.department_id):
        raise HTTPException(404, "Curriculum does not belong to the selected department")
    sections = get_section_options(db, request.academic_batch_id, term["semester_id"])
    section = next((s for s in sections if s["section_id"] == request.section_id), None)
    if section is None:
        raise HTTPException(404, "Section is not mapped to the selected curriculum and term")
    courses = get_course_options(db, request.academic_batch_id, term["semester_id"], section_id=request.section_id)
    if request.course_ids is not None:
        if not request.course_ids:
            raise HTTPException(422, "Select at least one course")
        if set(request.course_ids) - {c["course_id"] for c in courses}:
            raise HTTPException(422, "Selected courses are not mapped to this section and term")
        courses = [c for c in courses if c["course_id"] in request.course_ids]
    filters = {
        "department_id": request.department_id, "academic_batch_id": request.academic_batch_id,
        "academic_batch_name": batches[0]["academic_batch_desc"] or batches[0]["academic_batch_code"],
        "semester_id": term["semester_id"], "crclm_term_id": term["semester_id"],
        "semester_number": term["semester"], "term_name": term["semester_desc"] or f"Semester {term['semester']}",
        **section, "selected_course_ids": [c["course_id"] for c in courses],
        "include_total_marks": request.include_total_marks,
        "start_range": request.start_range, "end_range": request.end_range,
        "include_absents": request.include_absents,
        "marks_source": marks_source,
    }
    if not courses:
        return {"filters": filters, "rows": [], "courses": []}
    params = {"batch": request.academic_batch_id, "term": term["semester_id"],
              "section": request.section_id, "course_ids": filters["selected_course_ids"]}
    # A released question paper identifies an assessment, even when names repeat.
    assessments = _rows(db, """SELECT qp.qpd_id, qp.crs_id, qp.qpd_title,
        qp.qpd_max_marks, ao.ao_id, ao.ao_name, ao.max_marks, ao.section_id
        FROM cudos_qp_definition qp
        LEFT JOIN iems_assessment_occasions ao ON ao.qpd_id = qp.qpd_id
            AND ao.crs_id = qp.crs_id AND ao.academic_batch_id = qp.academic_batch_id
            AND ao.semester_id = qp.semester_id
        WHERE qp.academic_batch_id = :batch AND qp.semester_id = :term
          AND qp.crs_id IN :course_ids AND qp.qp_rollout > 1
          AND (ao.section_id = :section OR ao.mte_flag = 1 OR qp.qpd_type = 5)
        ORDER BY qp.crs_id, qp.qpd_id, ao.ao_id""", params, ("course_ids",))
    definitions = {}
    for a in assessments:
        key = (a["crs_id"], a["qpd_id"])
        # Prefer the selected section's description/max over a shared occasion.
        if key not in definitions or a["section_id"] == request.section_id:
            definitions[key] = {
                "component_id": str(a["qpd_id"]),
                "occasion_name": a["ao_name"] or a["qpd_title"] or f"Assessment {a['qpd_id']}",
                "max_marks": _number(a["max_marks"] if a["max_marks"] is not None else a["qpd_max_marks"]),
                "marks": None, "status": "missing", "source": "lms_assessment",
            }
    # Use explicit course enrollment for elective/batch membership. For regular
    # courses, the current student's curriculum and section form the roster.
    roster = _rows(db, """SELECT DISTINCT s.student_id, s.usno, s.regno, s.name,
        s.first_name, s.middle_name, s.last_name, c.crs_id
        FROM iems_students s
        JOIN iems_courses c ON c.crs_id IN :course_ids
        JOIN cudos_master_type_details sec ON sec.mt_details_id = :section
        WHERE (COALESCE(c.edu_sys_flag, 0) = 0 AND s.academic_batch_id = :batch
            AND s.section = sec.mt_details_name AND COALESCE(sec.parent_id, 0) = 0)
        OR EXISTS (SELECT 1 FROM cudos_map_courseto_student enrollment
            WHERE enrollment.student_id = s.student_id AND enrollment.crs_id = c.crs_id
              AND enrollment.semester_id = :term
              AND (enrollment.academic_batch_id = :batch OR s.academic_batch_id = :batch)
              AND ((COALESCE(sec.parent_id, 0) = 0 AND enrollment.section_id = :section)
                OR (COALESCE(sec.parent_id, 0) > 0 AND enrollment.batch_id = :section)))
        ORDER BY s.usno, s.student_id, c.crs_id""", params, ("course_ids",))
    marks = _rows(db, """SELECT sat.sat_id, sat.student_usn, sat.crs_id, sat.qpd_id,
        sat.total_marks FROM cudos_student_assessment_totalmarks sat
        JOIN cudos_qp_definition qp ON qp.qpd_id = sat.qpd_id AND qp.crs_id = sat.crs_id
        WHERE qp.academic_batch_id = :batch AND qp.semester_id = :term
          AND sat.crs_id IN :course_ids AND qp.qp_rollout > 1
          AND (sat.section_id = :section OR ((sat.section_id IS NULL OR sat.section_id = 0)
            AND (qp.qpd_type = 5 OR EXISTS (SELECT 1 FROM iems_assessment_occasions ao
                WHERE ao.qpd_id = qp.qpd_id AND ao.crs_id = qp.crs_id AND ao.mte_flag = 1))))
        ORDER BY sat.sat_id""", params, ("course_ids",))
    marks_by_student = {}
    for mark in marks:
        # Latest persisted total wins; duplicate totals must never be summed.
        marks_by_student[(mark["student_usn"], mark["crs_id"], str(mark["qpd_id"]))] = mark
    official_totals = {}
    if marks_source == "ems":
        definitions, marks, official_totals = _ems_data(db, params)
        marks_by_student = {(m["student_usn"], m["crs_id"], str(m["qpd_id"])): m for m in marks}
    students = {}
    membership = defaultdict(set)
    for s in roster:
        membership[s["student_id"]].add(s["crs_id"])
        students[s["student_id"]] = s
    headers = []
    for course in courses:
        headers.append({**course, "components": [dict(d) for (cid, _), d in definitions.items()
                                                  if cid == course["course_id"]]})
    rows = []
    ranged = request.start_range is not None
    for student_id, s in students.items():
        row_courses = []
        any_match = False
        for course in headers:
            components = []
            enrolled = course["course_id"] in membership[student_id]
            for definition in course["components"]:
                component = dict(definition)
                mark = marks_by_student.get((s["usno"], course["course_id"], component["component_id"])) if enrolled else None
                if mark is not None:
                    value = _number(mark["total_marks"])
                    matches = not ranged or (value is None and request.include_absents) or (
                        value is not None and request.start_range <= value <= request.end_range)
                    if matches:
                        component.update(marks=value, status="absent" if value is None else "marked")
                        any_match = True
                    else:
                        component["status"] = "filtered"
                components.append(component)
            numeric = [c["marks"] for c in components if c["marks"] is not None]
            total = round(sum(numeric), 2) if numeric else None
            if marks_source == "ems" and not ranged and enrolled:
                total = official_totals.get((s["usno"], course["course_id"]), total)
            row_courses.append({**course, "components": components,
                "total_marks": total,
                "data_available": enrolled and any(c["status"] in ("marked", "absent") for c in components)})
        if ranged and not any_match:
            continue
        name = s["name"] or " ".join(p for p in (s["first_name"], s["middle_name"], s["last_name"]) if p)
        rows.append({"sl_no": len(rows) + 1, "student_usn": s["usno"] or s["regno"] or str(student_id),
            "student_name": name or s["usno"] or str(student_id), "regno": s["regno"],
            "section": section["section_name"], "student_identity_status": "matched", "courses": row_courses})
    return {"filters": filters, "rows": rows, "courses": headers}


def build_consolidated_student_marks_graph(db, request, org_id=1):
    report = build_consolidated_student_marks_report(db, request, org_id)
    summaries = []
    for course in report["courses"]:
        enrolled = [c for r in report["rows"] for c in r["courses"] if c["course_id"] == course["course_id"]]
        totals = [c["total_marks"] for c in enrolled if c["total_marks"] is not None]
        assessments = []
        for definition in course["components"]:
            components = [component for c in enrolled for component in c["components"]
                          if component["component_id"] == definition["component_id"]]
            values = [c["marks"] for c in components if c["marks"] is not None]
            assessments.append({"component_id": definition["component_id"],
                "occasion_name": definition["occasion_name"], "max_marks": definition["max_marks"],
                "student_count": len(values), "absent_count": sum(c["status"] == "absent" for c in components),
                "average_marks": round(sum(values) / len(values), 2) if values else None})
        summaries.append({"course_id": course["course_id"], "course_code": course["course_code"],
            "course_title": course["course_title"], "student_count": len(totals),
            "average_marks": round(sum(totals) / len(totals), 2) if totals else None,
            "highest_marks": max(totals) if totals else None, "lowest_marks": min(totals) if totals else None,
            "assessments": assessments})
    return {"filters": report["filters"], "courses": summaries}


def export_consolidated_student_marks_report(db, request, org_id=1):
    report = build_consolidated_student_marks_report(db, request, org_id)
    headers = ["Sl. No", "USN", "Student Name"]
    columns = []
    for course in report["courses"]:
        for component in course["components"]:
            headers.append(f"{course['course_code']} - {component['occasion_name']} ({component['max_marks'] if component['max_marks'] is not None else '-'})")
            columns.append((course["course_id"], component["component_id"]))
        if request.include_total_marks:
            headers.append(f"{course['course_code']} - Total")
            columns.append((course["course_id"], None))
    data = []
    for student in report["rows"]:
        courses = {c["course_id"]: c for c in student["courses"]}
        row = [student["sl_no"], student["student_usn"], student["student_name"]]
        for cid, component_id in columns:
            course = courses[cid]
            if component_id is None:
                value = course["total_marks"]
            else:
                component = next(c for c in course["components"] if c["component_id"] == component_id)
                value = "AB" if component["status"] == "absent" else component["marks"]
            row.append("-" if value is None else value)
        data.append(row)
    # Treat user-origin text as literal text in spreadsheet applications.
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    if request.format == "csv":
        buffer = StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerows([[safe(v) for v in row] for row in [headers, *data]])
        return buffer.getvalue().encode("utf-8-sig"), "text/csv", "csv"
    if request.format == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Consolidated Marks"
        for row in [headers, *data]:
            sheet.append([safe(v) for v in row])
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24466B")
        for index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28 if index > 2 else 18
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    story = []
    # Split wide reports horizontally, repeating student identity on every page.
    for start in range(3, max(len(headers), 4), 6):
        indices = [0, 1, 2] + list(range(start, min(start + 6, len(headers))))
        if story:
            story.append(PageBreak())
        story.append(Paragraph("Consolidated Student Marks Report", styles["Heading2"]))
        story.append(Paragraph(escape(f"{report['filters']['academic_batch_name']} / {report['filters']['term_name']} / {report['filters']['section_name']}"), styles["Normal"]))
        cells = [[Paragraph(escape(str(row[i])), styles["Normal"]) for i in indices] for row in [headers, *data]]
        table = Table(cells, repeatRows=1, colWidths=[36, 88, 135] + [78] * (len(indices) - 3))
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), .3, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(table)
    SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=30, rightMargin=30).build(story)
    return buffer.getvalue(), "application/pdf", "pdf"
